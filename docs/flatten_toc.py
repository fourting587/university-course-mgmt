# -*- coding: utf-8 -*-
"""把 Word 更新后的 TOC 域转换为静态文本（保留页码与点线，去掉域底纹）。"""
import sys
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn

path = sys.argv[1]
doc = Document(path)

paras = doc.paragraphs
field_idx = None
entry_idx = []
for i, p in enumerate(paras):
    xml = p._p.xml
    if field_idx is None and "instrText" in xml and "TOC" in xml:
        field_idx = i
        continue
    if field_idx is not None and p.style is not None and p.style.name.lower().startswith("toc") and "\t" in p.text:
        entry_idx.append(i)

if field_idx is None or not entry_idx:
    print("no TOC field/entries found")
    sys.exit(1)

last = entry_idx[-1]
anchor = paras[last + 1] if last + 1 < len(paras) else None
if anchor is None:
    print("no anchor after TOC")
    sys.exit(1)

first = None
field_para = paras[field_idx]
if "\\t" in field_para.text:
    toc1 = next((p.style for p in paras if p.style and p.style.name.lower() == "toc 1"), field_para.style)
    first = (toc1, field_para.text.strip())
entries = ([] if first is None else [first]) + [(paras[i].style, paras[i].text) for i in entry_idx]
# 删除原域段落与目录段落
for i in [field_idx] + entry_idx:
    el = paras[i]._p
    el.getparent().remove(el)

for style, text in entries:
    p = anchor.insert_paragraph_before(style=style)
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.tab_stops.clear_all()
    p.paragraph_format.tab_stops.add_tab_stop(Cm(14.4), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    r = p.add_run(text)
    r.font.name = "宋体"
    r.font.size = Pt(12)
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")

# 去掉打开时自动更新域的设置
settings = doc.settings.element
for el in settings.findall(qn("w:updateFields")):
    settings.remove(el)

doc.save(path)
print("flattened", len(entries), "entries")
