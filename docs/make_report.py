# -*- coding: utf-8 -*-
"""基于实训报告模板生成完整报告。"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = r"C:\Users\86023\Desktop\university-course-mgmt"
TEMPLATE = os.path.join(BASE, "《信息化系统开发与应用)》实训报告模板.docx")
OUT = os.path.join(BASE, "docs", "信息化系统开发与应用实训报告-刘中天-20231026120.docx")
SHOT = os.path.join(BASE, "docs", "screenshots")

doc = Document(TEMPLATE)
# 去掉模板的文档网格，避免 1.5 倍行距被网格加倍
for _grid in doc.sections[0]._sectPr.findall(qn("w:docGrid")):
    doc.sections[0]._sectPr.remove(_grid)

# ---------- 字体与样式 ----------
def set_east_asia(style_or_run, name):
    rpr = style_or_run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)

normal = doc.styles["Normal"]
normal.font.name = "宋体"
normal.font.size = Pt(12)
set_east_asia(normal, "宋体")

def style(name, size, bold=False, east="宋体", color=None, before=0, after=0,
          line=None, indent=None, align=None, outline=None):
    st = doc.styles[name] if name in [s.name for s in doc.styles] else doc.styles.add_style(name, 1)
    st.font.name = east
    st.font.size = Pt(size)
    st.font.bold = bold
    if color:
        st.font.color.rgb = RGBColor(*color)
    set_east_asia(st, east)
    pf = st.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line
    if indent is not None:
        pf.first_line_indent = Pt(indent)
    if align is not None:
        pf.alignment = align
    if outline is not None:
        st.element.get_or_add_pPr().append(OxmlElement("w:outlineLvl"))
        st.element.pPr[-1].set(qn("w:val"), str(outline))
    return st

style("Heading 1", 15, True, "黑体", (0x0F, 0x3D, 0x4C), 18, 8, 1.5, None, WD_ALIGN_PARAGRAPH.LEFT, 0)
style("Heading 2", 13, True, "黑体", (0x1A, 0x5F, 0x7A), 12, 6, 1.5, None, WD_ALIGN_PARAGRAPH.LEFT, 1)
style("BodyCN", 12, False, "宋体", None, 0, 4, 1.5, 24, WD_ALIGN_PARAGRAPH.JUSTIFY)
style("CaptionCN", 10.5, False, "宋体", (0x47, 0x55, 0x69), 3, 8, 1.0, 0, WD_ALIGN_PARAGRAPH.CENTER)
style("CodeCN", 9, False, "Consolas", (0x1E, 0x29, 0x3B), 2, 6, 1.0, 0, WD_ALIGN_PARAGRAPH.LEFT)

for st in (doc.styles["Heading 1"], doc.styles["Heading 2"]):
    st.paragraph_format.page_break_before = False

def h1(text):
    p = doc.add_paragraph(text, style="Heading 1")
    p.paragraph_format.page_break_before = True
    return p

def h2(text):
    return doc.add_paragraph(text, style="Heading 2")

def body(text):
    for para in text.split("\n"):
        if para.strip():
            doc.add_paragraph(para.strip(), style="BodyCN")

def bullet(text):
    p = doc.add_paragraph(style="BodyCN")
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.left_indent = Pt(18)
    p.add_run("· " + text)
    return p

def caption(text):
    return doc.add_paragraph(text, style="CaptionCN")

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)

def table(headers, rows, widths=None, font_size=10.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        r = hdr[i].paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(font_size)
        r.font.name = "宋体"
        set_east_asia(r, "宋体")
        hdr[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade(hdr[i], "DCE6F1")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(font_size)
            r.font.name = "宋体"
            set_east_asia(r, "宋体")
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    return t

def image(path, width_cm, cap):
    doc.add_picture(path, width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(cap)

def code(text):
    p = doc.add_paragraph(style="CodeCN")
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F3F4F6")
    pPr.append(shd)
    p.add_run(text)
    return p

# ---------- 封面信息 ----------
cover = {6: ("实训选题：", "大学智能课程管理系统的设计与开发"),
         7: ("指导老师：", "廖景昌"),
         8: ("专业班级：", "23gb物联网工程1班"),
         9: ("学生姓名：", "刘中天"),
         10: ("学    号：", "20231026120")}
for idx, (label, value) in cover.items():
    p = doc.paragraphs[idx]
    if p.runs:
        p.runs[0].text = label + value
        for r in p.runs[1:]:
            r.text = ""

# 压缩封面行距，保证一页装下
for i in range(0, 11):
    p = doc.paragraphs[i]
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    size = 16 if i == 6 else 18
    if i in (0, 1, 2, 3, 5):
        size = 14
    for r in p.runs:
        r.font.size = Pt(size)

# ---------- 删除章节占位，保留封面；成绩表移到文末 ----------
table_el = doc.tables[0]._element
body_el = doc.element.body
body_el.remove(table_el)
for p in list(doc.paragraphs[11:]):
    p._p.getparent().remove(p._p)

# ---------- 页脚页码 ----------
sec = doc.sections[0]
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
def add_field(p, instr, placeholder=""):
    r = p.add_run()
    fld = OxmlElement("w:fldChar"); fld.set(qn("w:fldCharType"), "begin")
    r._r.append(fld)
    r2 = p.add_run()
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    r2._r.append(it)
    r3 = p.add_run()
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    r3._r.append(sep)
    r4 = p.add_run(placeholder)
    r4.font.size = Pt(9)
    r5 = p.add_run()
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    r5._r.append(end)
add_field(footer_p, " PAGE ", "1")

# ---------- 目录 ----------
from docx.enum.text import WD_BREAK

def page_break():
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    return p

page_break_removed = True
toc_title = doc.add_paragraph()
toc_title.paragraph_format.page_break_before = True
toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = toc_title.add_run("目    录")
tr.font.name = "黑体"
tr.font.size = Pt(18)
tr.bold = True
set_east_asia(tr, "黑体")
for _name in ("TOC 1", "TOC 2"):
    if _name not in [s.name for s in doc.styles]:
        _st = doc.styles.add_style(_name, 1)
    else:
        _st = doc.styles[_name]
    _st.font.name = "宋体"
    _st.font.size = Pt(12)
    set_east_asia(_st, "宋体")
    _st.paragraph_format.space_after = Pt(2)
    _st.paragraph_format.line_spacing = 1.2

toc_p = doc.add_paragraph()
add_field(toc_p, ' TOC \\o "1-2" \\z \\u ', "右键选择“更新域”生成目录")
for r in toc_p.runs:
    r.font.size = Pt(11)
    r.font.name = "宋体"
    set_east_asia(r, "宋体")

settings = doc.settings.element
update = OxmlElement("w:updateFields")
update.set(qn("w:val"), "true")
settings.append(update)

# ==================== 正文 ====================
h1("一．系统开发背景")

h2("1.1 实训背景")
body("""随着高校信息化建设的不断推进，课程管理已经从传统的纸质登记、Excel 统计逐步转向信息系统管理。课程管理涉及学生、教师、教务管理员三类角色，日常业务包括学生基本信息维护、班级与课程组织、考勤登记、课堂表现记录、实验与作业管理、平时成绩与总成绩汇总等。若完全依赖人工方式，容易出现数据重复、统计口径不一致、成绩计算错误、历史记录难以追溯等问题。""")
body("""本实训以大学日常教学课程管理为应用场景，设计并实现一套多角色、功能完善的智能课程管理系统。系统覆盖“基础数据维护—教学过程记录—成绩汇总—数据导出—系统审计”的完整链路，并采用前后端分离的部署形态，使系统既能满足课程演示，也具备一定的工程实践价值。""")

h2("1.2 实训目标与任务")
body("""本实训的总体目标是：综合运用 Java、Spring Boot、MyBatis、Thymeleaf、关系型数据库与容器化技术，完成一个可运行、可演示、可部署的课程管理系统，并在此过程中掌握需求分析、系统设计、编码实现、测试调试与部署发布的完整流程。具体任务包括：""")
bullet("完成开发环境搭建，统一使用 JDK 17、Maven、MySQL 8 与 Docker，保证编译与运行一致；")
bullet("完成需求分析，梳理学生、教师、管理员三类角色的功能边界与权限规则；")
bullet("完成系统设计，包括分层架构、功能模块、接口、数据库表结构与 E-R 模型；")
bullet("实现登录认证、考勤、课堂表现、实验、作业、成绩汇总、基础数据管理等核心功能；")
bullet("完成系统测试与缺陷修复，编写测试用例并验证关键业务规则；")
bullet("完成容器化部署，使用 Nginx 托管前端、Tomcat 部署后端 WAR、MySQL 存储数据。""")

h2("1.3 技术栈说明")
body("""系统后端采用 Java 17 与 Spring Boot 3.5.0，使用 Spring MVC 提供 Web 能力，MyBatis 3.0.5 负责数据访问，Thymeleaf 保留为服务端渲染界面，同时提供一套基于 JSON 的 REST 接口供独立前端调用。数据库采用 MySQL 8.0，字符集为 utf8mb4，存储引擎为 InnoDB，并通过外键约束保证参照完整性。前端采用原生 HTML、CSS 与 JavaScript 实现单页应用（SPA），通过 Nginx 托管静态资源并将 /api 请求反向代理到后端。系统使用 Docker Compose 编排 Nginx、Tomcat 与 MySQL 三个容器，实现一键部署。""")
body("""需要说明的是，课程指导手册中列出的数据库为 SQL Server，本系统实际部署环境采用 MySQL 8.0。两者均为关系型数据库，标准 SQL、事务、外键、索引等核心能力一致；本系统使用的 SQL 语法均控制在通用范围内，表结构、约束与索引设计可直接迁移到 SQL Server。选择 MySQL 主要是考虑课程机房与 Docker 环境的易获取性，以及 utf8mb4 对中文的友好支持。""")

h1("二．实验环境及配置介绍")
body("""本章介绍系统开发与运行所需的软件环境、安装配置步骤以及环境验证方法。实验环境以 Windows 11 为开发主机，使用 JDK 17 编译，MySQL 8.0 存储数据，Docker 负责运行期编排。""")

h2("2.1 开发工具与版本")
caption("表 2-1 开发环境与软件版本")
table(["类别", "软件/工具", "版本", "用途"], [
    ["操作系统", "Windows 11 64 位", "23H2", "开发与运行环境"],
    ["JDK", "Eclipse Temurin JDK", "17.0.x", "编译与运行 Java 程序"],
    ["构建工具", "Apache Maven", "3.9.x", "依赖管理与打包"],
    ["框架", "Spring Boot", "3.5.0", "Web、IoC、事务、嵌入式容器"],
    ["持久层", "MyBatis Spring Boot Starter", "3.0.5", "SQL 映射与数据访问"],
    ["模板引擎", "Thymeleaf", "3.1.x", "服务端渲染页面"],
    ["数据库", "MySQL Community Server", "8.0", "业务数据存储"],
    ["数据库客户端", "MySQL 命令行 / Navicat", "8.x", "建库、执行脚本、查询"],
    ["容器", "Docker Engine + Compose", "29.6 / v5.3", "镜像构建与多容器编排"],
    ["Web 服务器", "Nginx", "1.27", "静态资源托管与反向代理"],
    ["应用服务器", "Apache Tomcat", "10.1（JDK 17）", "部署 WAR 包"],
    ["浏览器", "Chrome / Edge", "最新版", "系统访问与调试"],
], widths=[3.0, 4.6, 3.2, 4.2], font_size=10)

h2("2.2 环境安装与配置步骤")
body("""（1）安装 JDK 17。双击安装包完成安装后，配置 JAVA_HOME 指向 JDK 17 安装目录，并将 %JAVA_HOME%\\bin 加入 Path。执行 java -version 与 javac -version，确认输出均为 17。若 IDE 中同时安装了更高版本 JDK，需要将项目的编译 JDK 固定为 17，否则会出现 class 文件版本 69.0 无法加载的问题。""")
body("""（2）安装并配置 Maven。解压 Maven 后配置 MAVEN_HOME 与 Path，执行 mvn -v 检查版本。为提升依赖下载速度，可在 settings.xml 中配置国内镜像仓库。""")
body("""（3）安装 MySQL 8.0。安装完成后启动 MySQL 服务，使用 root 账号登录，执行建库语句：CREATE DATABASE course_mgmt CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci; 随后依次执行 schema-mysql.sql 与 data-mysql.sql，完成建表、外键创建与演示数据导入。""")
body("""（4）导入项目并启动。使用 IntelliJ IDEA 或 VS Code 打开项目，等待 Maven 依赖下载完成，运行 CmsApplication 启动类；也可以执行 mvn clean package 生成可执行 WAR，再通过 java -jar 运行。启动成功后访问 http://localhost:1001 即可看到登录页面。""")
body("""（5）容器化运行。安装并启动 Docker Desktop 后，在项目根目录执行 docker compose up -d --build，Compose 会依次启动 MySQL、Tomcat 与 Nginx 三个容器。MySQL 首次启动时自动执行建表与数据脚本，Tomcat 部署 ROOT.war，Nginx 托管前端静态资源并反代 /api。启动完成后访问 http://localhost:8080 进入前后端分离版界面。""")

h2("2.3 环境测试与运行结果")
body("""环境搭建完成后进行了以下验证：java -version 输出 17；mvn -v 输出 Maven 3.9；mysql -uroot -p 能正常连接并看到 course_mgmt 数据库；mvn clean package 构建成功并生成 course-management-system-1.0.0.war；docker compose ps 显示 cms-web、cms-app、cms-mysql 三个容器均为 Up 状态，其中 MySQL 健康检查通过。浏览器访问 8080 端口可正常打开登录页，使用 admin/123456 登录后按提示修改密码即可进入工作台。""")
body("""构建过程中曾出现两个环境问题：一是 IDE 使用 JDK 25 编译导致运行时报 UnsupportedClassVersionError（class 文件版本 69.0），通过将 IDE 的编译 JDK 固定为 17 并执行 mvn clean package 解决；二是重复执行数据脚本导致学生数据重复，通过移除启动时的数据脚本自动执行、仅在初始化时手工执行一次快照解决。""")

h1("三．系统需求")
body("""本章从角色、功能与非功能三个维度梳理系统需求，并给出成绩计算等关键业务规则。需求分析遵循“角色—场景—功能—数据”的思路，确保每个功能都能落到具体的页面、接口与数据表。""")

h2("3.1 角色与用例")
caption("表 3-1 角色与主要用例")
table(["角色", "主要用例", "数据范围"], [
    ["学生", "登录、修改密码、查询本人考勤/表现/实验/作业、查看平时与总成绩", "仅本人数据"],
    ["教师", "登录、查看所授课程学生名单、考勤单条/批量录入、表现录入、实验与作业管理、成绩汇总与导出、期末成绩录入", "本人所授课程"],
    ["管理员", "学生/教师/班级/课程/评价规则 CRUD、重置密码、启用禁用、全部教学数据查询、操作日志", "全校数据"],
], widths=[2.6, 10.0, 3.0], font_size=10)

h2("3.2 功能性需求")
body("""（1）登录认证与权限控制。系统需区分学生、教师、管理员三种角色，登录成功后根据角色展示不同菜单；未登录用户访问业务页面时应跳转登录页；跨角色访问应被拒绝；使用默认密码登录后必须强制修改密码；连续多次登录失败应临时锁定账号。""")
body("""（2）考勤管理。教师与管理员可按课程、班级、日期录入考勤，支持单条录入与整班批量录入；同一学生、同一课程、同一日期只能存在一条记录；支持按课程、班级、学期筛选与分页查看；学生只能查看本人考勤。""")
body("""（3）课堂表现。教师可按课程与学生录入表现得分与评语，支持查询、删除；得分范围限定为 0—100。""")
body("""（4）实验管理。教师可创建实验（名称、满分、截止日期），按学生录入完成度、得分与备注；支持查询、修改与删除；完成度与得分均需校验范围。""")
body("""（5）作业管理。教师可创建作业（名称、满分、截止日期），按学生录入提交状态、得分与备注；支持查询、修改与删除。""")
body("""（6）成绩统计。系统按评价规则将考勤、课堂表现、实验、作业加权得到平时成绩，再按平时占比与期末占比计算总成绩；支持按课程、班级筛选，分页展示，并可导出 CSV 由 Excel 打开。""")
body("""（7）基础数据管理。管理员可对学生、教师、班级、课程、评价规则进行增删改查；删除班级前需先清空学生，删除教师前需先转移课程；删除课程时级联清理其考勤、表现、实验、作业、期末成绩与评价规则。""")
body("""（8）操作审计。系统记录登录成功/失败、密码修改、基础数据增删改、教学数据录入与删除、成绩导出等关键操作，管理员可按时间倒序分页查看。""")

h2("3.3 非功能性需求")
body("""易用性：页面布局统一，左侧导航按角色收缩；表单字段带默认值与范围限制；删除操作二次确认；列表分页避免一次性加载大量数据；错误信息以中文提示。""")
body("""稳定性：数据库使用 InnoDB 事务与外键约束；写操作使用事务保证一致性；分页查询避免全表渲染；容器配置健康检查与自动重启；应用异常统一返回友好错误页或 JSON。""")
body("""安全性：密码使用 BCrypt 加盐哈希存储，兼容历史明文并自动升级；使用 Session 维持登录态，改密后使其他设备会话失效；POST 请求校验 Origin/Referer 同源；接口按角色与课程归属双重校验；登录失败锁定降低暴力破解风险。""")

h2("3.4 关键业务规则")
body("""（1）考勤折算：到课 100 分、迟到 80 分、请假 60 分、缺勤 0 分，按记录条数取平均得到考勤分。""")
body("""（2）平时成绩 = 考勤分×考勤权重 + 表现分×表现权重 + 实验分×实验权重 + 作业分×作业权重，权重之和为 100。""")
body("""（3）总成绩 = 平时成绩×平时占比 + 期末成绩×期末占比，平时占比与期末占比之和为 100。""")
body("""（4）评价规则的六项数值均需在 0—100 之间；录入成绩、完成度时统一校验 0—100 范围，实验与作业满分允许 0—1000。""")

h1("四．系统概要设计")
body("""本章介绍系统的总体架构、功能模块划分、接口设计、界面设计与关键业务流程。系统采用前后端分离的部署形态，同时保留 Thymeleaf 服务端渲染页面作为兼容界面。""")

h2("4.1 总体架构")
image(os.path.join(BASE, "docs", "部署架构图.png"), 8.5, "图 4-1 系统部署架构图")
body("""系统分为表示层、接口层、业务层、数据访问层与数据层。表示层由独立前端（原生 JS 单页应用）与 Thymeleaf 页面组成；接口层由 Spring MVC 的 Controller 提供页面路由与 /api JSON 接口；业务层由 Service 承担成绩计算、基础数据维护、审计等业务逻辑；数据访问层由 MyBatis Mapper 通过注解 SQL 访问数据库；数据层为 MySQL 8.0。运行期由 Nginx 托管前端并反向代理 /api 到 Tomcat，Tomcat 以 ROOT.war 形式部署后端，MySQL 独立容器存储数据。""")
body("""分层设计带来的好处是：页面与接口解耦，前端可以独立迭代；Controller 只负责参数接收与视图/JSON 返回，业务规则集中在 Service；Mapper 只关注 SQL，便于优化与复用。系统还在 Controller 与 Mapper 之间保留了 Service 扩展点，后续可逐步将各教学模块的业务逻辑上移，进一步提升可维护性。""")

h2("4.2 功能模块设计")
image(os.path.join(BASE, "docs", "功能模块图.png"), 15.0, "图 4-2 系统功能模块图")
body("""系统功能按角色划分为登录认证、学生端、教师端、管理员端与公共支撑五个部分。登录认证负责账号密码校验、角色识别、默认密码强制修改与失败锁定；学生端提供本人数据查询；教师端提供教学过程数据录入与成绩管理；管理员端提供基础数据维护、账号管理与审计；公共支撑包括 REST 接口、容器化部署与数据备份。""")

h2("4.3 接口设计")
body("""前后端分离版以 /api 为统一前缀，采用 REST 风格，请求与响应均为 JSON（导出接口返回 CSV）。未登录返回 401，默认密码未修改返回 403 与 MUST_CHANGE 标识，权限不足返回 403。主要接口如表 4-1 所示。""")
caption("表 4-1 主要接口设计")
table(["方法", "路径", "说明", "权限"], [
    ["POST", "/api/auth/login", "登录，返回用户与是否需改密", "公开"],
    ["GET", "/api/auth/me", "当前登录用户", "登录"],
    ["POST", "/api/auth/password", "修改密码", "登录"],
    ["GET", "/api/dashboard", "工作台统计", "登录"],
    ["GET", "/api/courses", "按角色返回课程列表", "登录"],
    ["GET", "/api/students", "学生分页查询", "管理员"],
    ["GET", "/api/attendance", "考勤分页查询", "登录（按角色收窄）"],
    ["POST", "/api/attendance", "单条考勤录入", "教师/管理员"],
    ["POST", "/api/attendance/batch", "批量考勤录入", "教师/管理员"],
    ["POST", "/api/performance", "课堂表现录入", "教师/管理员"],
    ["POST", "/api/experiments", "新建/修改实验", "教师/管理员"],
    ["POST", "/api/experiments/records", "实验记录录入", "教师/管理员"],
    ["POST", "/api/homework", "新建/修改作业", "教师/管理员"],
    ["POST", "/api/exams", "期末成绩录入（自动覆盖）", "教师/管理员"],
    ["GET", "/api/scores", "成绩报表分页", "登录"],
    ["GET", "/api/scores/export", "导出 CSV", "登录"],
    ["GET", "/api/audit", "操作日志分页", "管理员"],
], widths=[2.2, 7.0, 5.6, 3.4], font_size=9.5)

h2("4.4 界面设计")
body("""系统提供两套界面：8080 端口的前后端分离版与 1001 端口的 Thymeleaf 版。两套界面共用同一套数据库与权限规则。前后端分离版采用左侧导航 + 顶部栏 + 内容卡片的布局，表格统一分页，录入操作通过弹窗表单完成，删除操作二次确认。主要页面如表 4-2 所示。""")
caption("表 4-2 主要页面设计")
table(["页面", "路径", "主要内容"], [
    ["登录", "/（前端路由）", "账号密码登录、默认密码提示"],
    ["工作台", "#/dashboard", "学生/教师/课程/班级/考勤数量统计"],
    ["成绩汇总", "#/scores", "课程与类型筛选、分页、导出"],
    ["考勤", "#/attendance", "筛选、单条/批量录入、删除"],
    ["课堂表现", "#/performance", "筛选、录入、删除"],
    ["实验", "#/experiments", "新建实验、录入记录、删除"],
    ["作业", "#/homework", "新建作业、录入记录、删除"],
    ["期末录入", "#/exam", "按课程列出学生并录入/修改期末成绩"],
    ["学生管理", "#/students", "查询、新增、编辑、删除、重置密码"],
    ["班级/教师/课程/规则", "#/classes 等", "基础数据增删改查"],
    ["操作日志", "#/audit", "审计记录分页查看"],
], widths=[4.0, 5.2, 8.0], font_size=10)

h2("4.5 关键业务流程")
body("""登录流程：浏览器提交账号密码 → Controller 校验锁定状态 → 查询用户并校验 BCrypt 密码 → 判断是否默认密码 → 写入 Session 并注册会话 → 默认密码跳转改密页，否则进入工作台。""")
body("""成绩计算流程：选择课程与班级 → 权限校验 → 查询课程绑定的班级与学生 → 一次性聚合查询考勤、表现、实验、作业、期末数据 → 按评价规则加权 → 分页返回。聚合查询将原先按学生逐条查询的 N+1 问题优化为每类数据一次查询。""")
body("""考勤批量流程：选择课程与日期 → 加载课程学生名单 → 为每名学生选择状态 → 一次性提交记录数组 → 后端逐条插入并统计成功/跳过数量 → 返回结果并刷新列表。""")

h1("五．系统 E-R 图设计")
body("""系统数据库共 16 张表，按业务域可分为基础数据与教学数据两组。为避免单张 E-R 图过于拥挤，本文将其拆分为基础数据 E-R 图与教学数据 E-R 图，分别描述实体、属性及联系。""")

h2("5.1 实体与属性")
body("""基础数据实体包括：角色（sys_role）、用户（sys_user）、班级（clazz）、学生（student）、教师（teacher）、课程（course）、选课关系（course_class）与评价规则（eval_rule）。教学数据实体包括：学生、课程、考勤（attendance）、课堂表现（performance）、期末成绩（exam_score）、实验（experiment）、实验记录（experiment_record）、作业（homework）与作业记录（homework_record）。此外，audit_log 为独立的审计实体，不与其他表建立外键。""")

h2("5.2 基础数据 E-R 图")
image(os.path.join(BASE, "docs", "ER图-基础数据.png"), 11.0, "图 5-1 基础数据 E-R 图")
body("""基础数据 E-R 图如图 5-1 所示。角色与用户为一对多关系；班级与学生为一对多关系；教师与课程为一对多关系；课程与班级通过 course_class 形成多对多关系；课程与评价规则为一对一关系。""")

h2("5.3 教学数据 E-R 图")
image(os.path.join(BASE, "docs", "ER图-教学数据.png"), 12.5, "图 5-2 教学数据 E-R 图")
body("""教学数据 E-R 图如图 5-2 所示。学生与考勤、表现、期末成绩、实验记录、作业记录均为一对多关系；课程与考勤、表现、期末成绩、实验、作业均为一对多关系；实验与实验记录、作业与作业记录为一对多关系。""")

h2("5.4 关系说明")
body("""（1）一个班级包含多名学生，一名学生只属于一个班级，学生表的 clazz_id 为外键。""")
body("""（2）一名教师可以教授多门课程，一门课程只由一名教师负责，课程表的 teacher_id 为外键。""")
body("""（3）一门课程可以面向多个班级开课，一个班级可以有多门课程，二者通过 course_class 关联，联合唯一约束避免重复选课。""")
body("""（4）一条考勤记录唯一对应一名学生、一门课程与一个日期，三列联合唯一；表现、期末成绩同理受学生与课程约束。""")
body("""（5）实验记录与作业记录以“实验/作业 + 学生”联合唯一，重复录入时执行更新而非新增。""")
body("""（6）评价规则与课程一对一，用于保存考勤、表现、实验、作业权重及平时、期末占比。""")

h1("六．系统数据库设计")
body("""数据库设计遵循第三范式，兼顾查询性能与数据完整性。所有表使用 InnoDB 引擎与 utf8mb4 字符集，主键统一为自增 INT，时间字段使用 TIMESTAMP，业务外键均建立索引。""")

h2("6.1 设计原则")
body("""（1）规范化：基础数据与业务数据分离，避免冗余；学生只保存班级外键，班级名称通过关联查询获取。""")
body("""（2）完整性：通过外键约束保证参照完整性；学生、教师被引用时使用 RESTRICT，教学数据使用 CASCADE，删除课程时自动清理相关记录。""")
body("""（3）一致性：考勤、实验记录、作业记录、期末成绩均设置联合唯一约束，防止重复录入；评价规则课程唯一。""")
body("""（4）可扩展：预留 audit_log 审计表，记录关键操作；用户表通过 role_code 与 ref_id 关联学生/教师，支持后续扩展更多角色。""")

h2("6.2 数据表清单")
caption("表 6-1 数据表清单")
table(["序号", "表名", "说明", "关键字段"], [
    ["1", "sys_role", "角色表", "role_code、role_name"],
    ["2", "sys_user", "用户账号表", "username、password、role_code、ref_id"],
    ["3", "clazz", "班级表", "class_no、class_name、grade、major"],
    ["4", "student", "学生表", "student_no、name、clazz_id"],
    ["5", "teacher", "教师表", "teacher_no、name、title"],
    ["6", "course", "课程表", "course_no、course_name、credit、teacher_id"],
    ["7", "course_class", "课程-班级关联表", "course_id、clazz_id"],
    ["8", "eval_rule", "评价规则表", "course_id、四项权重、两项占比"],
    ["9", "attendance", "考勤表", "student_id、course_id、attend_date、status"],
    ["10", "performance", "课堂表现表", "student_id、course_id、record_date、score"],
    ["11", "experiment", "实验表", "course_id、title、full_score、due_date"],
    ["12", "experiment_record", "实验记录表", "experiment_id、student_id、completion、score"],
    ["13", "homework", "作业表", "course_id、title、full_score、due_date"],
    ["14", "homework_record", "作业记录表", "homework_id、student_id、status、score"],
    ["15", "exam_score", "期末成绩表", "student_id、course_id、score、exam_date"],
    ["16", "audit_log", "操作审计表", "username、action、target、ip、created_at"],
], widths=[1.4, 4.0, 4.6, 7.4], font_size=9.5)

h2("6.3 核心表结构")
caption("表 6-2 学生表 student")
table(["字段", "类型", "约束", "说明"], [
    ["id", "INT", "PK，自增", "学生主键"],
    ["student_no", "VARCHAR(30)", "NOT NULL，UNIQUE", "学号"],
    ["name", "VARCHAR(50)", "NOT NULL", "姓名"],
    ["gender", "VARCHAR(4)", "可空", "性别"],
    ["clazz_id", "INT", "NOT NULL，FK→clazz", "所属班级"],
    ["phone", "VARCHAR(20)", "可空", "联系电话"],
    ["email", "VARCHAR(80)", "可空", "邮箱"],
], widths=[3.4, 3.6, 5.0, 4.0], font_size=9.5)

caption("表 6-3 考勤表 attendance")
table(["字段", "类型", "约束", "说明"], [
    ["id", "INT", "PK，自增", "考勤主键"],
    ["student_id", "INT", "NOT NULL，FK→student", "学生"],
    ["course_id", "INT", "NOT NULL，FK→course", "课程"],
    ["attend_date", "DATE", "NOT NULL", "上课日期"],
    ["status", "VARCHAR(20)", "NOT NULL", "PRESENT/LATE/LEAVE/ABSENT"],
    ["remark", "VARCHAR(200)", "可空", "备注"],
    ["recorder_id", "INT", "可空", "录入人用户 id"],
    ["(student_id,course_id,attend_date)", "联合唯一", "UNIQUE", "防止重复考勤"],
], widths=[5.6, 3.2, 4.2, 4.0], font_size=9.5)

caption("表 6-4 评价规则表 eval_rule")
table(["字段", "类型", "约束", "说明"], [
    ["id", "INT", "PK，自增", "规则主键"],
    ["course_id", "INT", "NOT NULL，UNIQUE，FK→course", "课程"],
    ["attendance_weight", "DECIMAL(5,2)", "默认 20", "考勤权重"],
    ["performance_weight", "DECIMAL(5,2)", "默认 20", "表现权重"],
    ["experiment_weight", "DECIMAL(5,2)", "默认 30", "实验权重"],
    ["homework_weight", "DECIMAL(5,2)", "默认 30", "作业权重"],
    ["regular_ratio", "DECIMAL(5,2)", "默认 40", "平时占比"],
    ["exam_ratio", "DECIMAL(5,2)", "默认 60", "期末占比"],
], widths=[4.4, 3.6, 5.4, 3.6], font_size=9.5)

caption("表 6-5 实验记录表 experiment_record")
table(["字段", "类型", "约束", "说明"], [
    ["id", "INT", "PK，自增", "记录主键"],
    ["experiment_id", "INT", "NOT NULL，FK→experiment", "实验"],
    ["student_id", "INT", "NOT NULL，FK→student", "学生"],
    ["score", "DECIMAL(5,2)", "可空", "实验得分"],
    ["completion", "DECIMAL(5,2)", "NOT NULL，默认 0", "完成度"],
    ["submit_time", "TIMESTAMP", "可空", "提交时间"],
    ["remark", "VARCHAR(200)", "可空", "备注"],
    ["(experiment_id,student_id)", "联合唯一", "UNIQUE", "每人每实验一条"],
], widths=[5.4, 3.2, 4.4, 3.8], font_size=9.5)

h2("6.4 外键与级联策略")
caption("表 6-6 课程表 course")
table(["字段", "类型", "约束", "说明"], [
    ["id", "INT", "PK，自增", "课程主键"],
    ["course_no", "VARCHAR(30)", "NOT NULL，UNIQUE", "课程号"],
    ["course_name", "VARCHAR(100)", "NOT NULL", "课程名称"],
    ["credit", "DECIMAL(3,1)", "默认 2.0", "学分"],
    ["semester", "VARCHAR(20)", "NOT NULL", "开课学期"],
    ["teacher_id", "INT", "NOT NULL，FK→teacher", "任课教师"],
], widths=[3.4, 3.6, 5.0, 4.0], font_size=9.5)

caption("表 6-7 作业记录表 homework_record")
table(["字段", "类型", "约束", "说明"], [
    ["id", "INT", "PK，自增", "记录主键"],
    ["homework_id", "INT", "NOT NULL，FK→homework", "作业"],
    ["student_id", "INT", "NOT NULL，FK→student", "学生"],
    ["score", "DECIMAL(5,2)", "可空", "作业得分"],
    ["status", "VARCHAR(20)", "默认 PENDING", "SUBMITTED/GRADED/PENDING"],
    ["submit_time", "TIMESTAMP", "可空", "提交时间"],
    ["remark", "VARCHAR(200)", "可空", "备注"],
    ["(homework_id,student_id)", "联合唯一", "UNIQUE", "每人每作业一条"],
], widths=[5.4, 3.2, 4.6, 3.6], font_size=9.5)

caption("表 6-8 用户表 sys_user")
table(["字段", "类型", "约束", "说明"], [
    ["id", "INT", "PK，自增", "用户主键"],
    ["username", "VARCHAR(50)", "NOT NULL，UNIQUE", "登录名（学号/工号）"],
    ["password", "VARCHAR(100)", "NOT NULL", "BCrypt 密文"],
    ["role_code", "VARCHAR(20)", "NOT NULL", "ADMIN/TEACHER/STUDENT"],
    ["ref_id", "INT", "可空", "关联学生/教师 id"],
    ["real_name", "VARCHAR(50)", "NOT NULL", "姓名"],
    ["enabled", "BOOLEAN", "默认 TRUE", "是否启用"],
], widths=[3.4, 3.8, 4.8, 4.0], font_size=9.5)

caption("表 6-9 主要外键与删除策略")
table(["外键", "引用", "删除策略", "原因"], [
    ["student.clazz_id", "clazz.id", "RESTRICT", "班级有学生时不允许删除"],
    ["course.teacher_id", "teacher.id", "RESTRICT", "教师有课程时不允许删除"],
    ["course_class.course_id", "course.id", "CASCADE", "删课程时清理选课关系"],
    ["course_class.clazz_id", "clazz.id", "CASCADE", "删班级时清理选课关系"],
    ["eval_rule.course_id", "course.id", "CASCADE", "规则随课程删除"],
    ["attendance.student_id", "student.id", "CASCADE", "学生删除后清理考勤"],
    ["attendance.course_id", "course.id", "CASCADE", "课程删除后清理考勤"],
    ["performance/exam_score", "student/course", "CASCADE", "同上"],
    ["experiment_record.experiment_id", "experiment.id", "CASCADE", "实验删除后清理记录"],
    ["homework_record.homework_id", "homework.id", "CASCADE", "作业删除后清理记录"],
], widths=[5.4, 3.6, 3.0, 4.8], font_size=9.5)

h2("6.5 数据初始化与快照")
body("""为便于演示与验收，系统提供完整的数据库快照脚本 data-mysql.sql。脚本先关闭外键检查并 TRUNCATE 全部业务表，再按依赖顺序批量插入数据，最后恢复外键检查。快照包含 8 个班级、240 名学生、5 名教师、12 门课程、76 条课程班级关系、12 条评价规则，以及按 2025 年秋季学期 16 个教学周生成的约 6.5 万条考勤、表现、实验、作业与期末记录。所有上课日期均为工作日，避开国庆假期，实验与作业提交时间落在工作日晚上，保证演示数据的真实性。""")

h1("七．系统功能实现")
body("""本章介绍核心功能的实现思路与关键代码。系统后端采用 Controller—Service—Mapper 分层，Controller 负责参数接收与权限入口，Service 承担业务规则，Mapper 通过注解 SQL 访问数据库。""")

h2("7.1 登录认证与权限控制")
body("""登录接口先检查账号是否因连续失败被锁定，再查询用户并使用 BCrypt 校验密码。校验通过后判断密码是否为默认值 123456，若是则在 Session 中写入强制改密标记，后续请求只允许访问改密接口。登录成功后注册会话，改密时使该账号在其他设备上的会话失效。""")
body("""权限控制由 LoginInterceptor 统一完成：未登录访问页面跳转登录页，访问 /api 返回 401；强制改密状态下除改密与退出外的请求一律拦截；/admin 仅管理员可访问，/teach 仅教师或管理员可访问。对于课程级数据，Controller 还通过 CourseAccess 校验课程是否属于当前教师，避免横向越权。""")
body("""此外，系统对 POST 请求校验 Origin/Referer 与当前站点是否同源，防止跨站请求伪造；登录失败 5 次锁定 10 分钟；密码使用 BCrypt 加盐哈希，历史明文密码在首次成功登录后自动升级为密文。""")

h2("7.2 考勤管理")
body("""考勤支持单条与批量两种录入方式。单条录入选择课程、学生、日期与状态，后端校验课程归属后插入记录；若同一学生、课程、日期已存在记录，则返回冲突提示。批量录入先按课程加载学生名单，前端为每名学生选择状态，提交时组装记录数组，后端循环插入并统计成功与跳过数量。""")
body("""考勤查询支持按课程、班级、学期筛选并分页，学生账号会被强制附加本人 student_id 条件，教师账号会被限制在其所授课程范围内。""")

h2("7.3 课堂表现")
body("""课堂表现记录包含学生、课程、日期、得分与评语。后端对得分执行 0—100 校验，教师只能操作本人课程；列表按课程筛选并分页，支持删除。""")

h2("7.4 实验管理")
body("""实验管理分为实验定义与实验记录两部分。实验定义包含课程、名称、满分与截止日期；实验记录按学生录入完成度、得分与备注。后端校验完成度与得分范围，并在重复录入同一学生同一实验时自动转为更新，避免联合唯一约束报错。""")

h2("7.5 作业管理")
body("""作业管理与实验类似，包含作业定义与作业记录。作业记录保存提交状态、得分、提交时间与备注，状态支持已提交、已批改、未交；重复录入同样自动更新。""")

h2("7.6 成绩统计")
body("""成绩统计是系统的核心业务。系统先根据课程评价规则读取四项权重与两项占比，再聚合每名学生的考勤、表现、实验、作业与期末数据：考勤按状态折算为分数后取平均，表现、实验、作业取平均分，期末取录入成绩。随后计算平时成绩与总成绩，并按课程、班级筛选、分页展示。""")
body("""为避免 N+1 查询，ScoreService 先一次性查询每类数据的聚合结果并放入 Map，再按学生匹配，使每类数据只访问数据库一次。成绩报表支持导出 CSV，文件带 UTF-8 BOM，Excel 打开中文不乱码。""")
code("""// ScoreService：一次性聚合，避免按学生循环查询
Map<Integer, Double> attendance = toMap(attendanceMapper.countByCourse(courseId));
Map<Integer, Double> performance = toMap(performanceMapper.avgByCourse(courseId));
Map<Integer, Double> experiment = toMap(experimentMapper.avgByCourse(courseId));
Map<Integer, Double> homework = toMap(homeworkMapper.avgByCourse(courseId));
Map<Integer, Double> exam = toMap(examScoreMapper.scoreByCourse(courseId));

double regular = attendanceScore * w.attendanceWeight
        + performanceScore * w.performanceWeight
        + experimentScore * w.experimentWeight
        + homeworkScore * w.homeworkWeight;
regular = regular / 100.0;
double total = regular * regularRatio / 100.0 + examScore * examRatio / 100.0;""")

h2("7.7 基础数据管理")
body("""管理员可对学生、教师、班级、课程与评价规则进行增删改查。新增学生或教师时同步创建登录账号，默认密码为 123456，修改姓名时同步更新账号显示名；删除学生或教师时同步删除账号。删除班级前检查是否仍有学生，删除教师前检查是否仍有课程；删除课程时级联清理选课关系、评价规则、考勤、表现、实验、作业与期末成绩。""")
body("""评价规则保存时校验六项数值均在 0—100 之间，避免出现负权重或超过 100 的占比。所有管理操作写入审计日志。""")

h2("7.8 前后端分离与部署")
body("""系统提供 /api REST 接口与独立前端。前端为零依赖的原生 JS 单页应用，使用 hash 路由与 fetch 调用接口，登录态通过同源 Cookie 维持；Nginx 托管静态资源并将 /api 反向代理到后端，因此不存在跨域 Cookie 问题。后端打包为可执行 WAR，既可直接 java -jar 运行，也可部署到外部 Tomcat。""")
body("""Docker Compose 编排三个服务：cms-mysql 首次启动执行建表与快照脚本；cms-app 基于 tomcat:10.1-jdk17-temurin 镜像部署 ROOT.war；cms-web 基于 nginx:1.27 托管前端并反代 /api。Compose 通过健康检查保证应用在 MySQL 就绪后启动。""")

h2("7.9 关键代码")
body("""（1）登录拦截器核心逻辑：统一处理未登录、强制改密、角色越权与 API 的 JSON 响应。""")
code("""public boolean preHandle(HttpServletRequest request, HttpServletResponse response, Object handler) {
    HttpSession session = request.getSession();
    LoginUser user = (LoginUser) session.getAttribute(Constants.SESSION_USER);
    String uri = request.getRequestURI();

    if ("POST".equalsIgnoreCase(request.getMethod()) && !isSameOrigin(request)) {
        response.sendError(403, "Cross-site request blocked");
        return false;
    }
    if (user == null) {
        if (uri.startsWith("/api/")) return writeJson(response, 401, "{\\"error\\":\\"未登录\\"}");
        response.sendRedirect(request.getContextPath() + "/login");
        return false;
    }
    if (Boolean.TRUE.equals(session.getAttribute(Constants.SESSION_MUST_CHANGE))
            && !uri.startsWith("/password/change") && !uri.startsWith("/api/auth/")) {
        if (uri.startsWith("/api/")) return writeJson(response, 403, "{\\"error\\":\\"请先修改默认密码\\"}");
        response.sendRedirect(request.getContextPath() + "/password/change");
        return false;
    }
    return true;
}""")
body("""（2）批量考勤：同一课程、同一日期逐条写入，重复记录跳过并计数。""")
code("""@PostMapping("/attendance/batch")
public ResponseEntity<?> batchAttendance(@RequestBody Map<String, Object> body, HttpSession session) {
    Integer courseId = intg(body, "courseId");
    if (!courseAccess.canAccess(user(session), courseId)) return forbidden();
    LocalDate day = LocalDate.parse(str(body, "attendDate"));
    int saved = 0, skipped = 0;
    for (Map<String, Object> item : records) {
        Attendance a = new Attendance();
        a.setCourseId(courseId);
        a.setStudentId(intg(item, "studentId"));
        a.setAttendDate(java.sql.Date.valueOf(day));
        a.setStatus(str(item, "status"));
        try { attendanceMapper.insert(a); saved++; }
        catch (Exception e) { skipped++; }
    }
    return ResponseEntity.ok(Map.of("ok", true, "saved", saved, "skipped", skipped));
}""")

h2("7.10 系统测试与运行结果")
body("""系统测试采用功能测试为主、接口测试与单元测试为辅的方式。功能测试覆盖登录、权限、考勤、表现、实验、作业、成绩、基础数据与导出；接口测试使用 HTTP 请求验证状态码与返回结构；单元测试覆盖密码工具、登录锁定、课程权限、登录拦截器与成绩计算。""")
caption("表 7-1 核心测试用例（完整 30 条见附件《测试用例.xlsx》）")
table(["编号", "测试项", "输入/操作", "预期结果", "结果"], [
    ["TC-01", "正确登录", "admin/123456", "提示修改默认密码", "通过"],
    ["TC-02", "错误密码", "admin/000000", "提示用户名或密码错误", "通过"],
    ["TC-03", "登录锁定", "连续 5 次错误密码", "锁定 10 分钟", "通过"],
    ["TC-04", "未登录访问 API", "直接请求 /api/dashboard", "401 JSON", "通过"],
    ["TC-05", "学生越权", "学生访问 /api/admin/students", "403", "通过"],
    ["TC-06", "教师跨课程", "教师给非本人课程录考勤", "403", "通过"],
    ["TC-07", "单条考勤", "选择课程/学生/日期", "新增成功", "通过"],
    ["TC-08", "重复考勤", "同一学生课程日期再次提交", "返回冲突提示", "通过"],
    ["TC-09", "批量考勤", "整班提交 30 条", "保存成功并统计", "通过"],
    ["TC-10", "成绩范围", "表现分录入 120", "400 并提示 0—100", "通过"],
    ["TC-11", "权重校验", "规则权重录入 150", "400 并提示 0—100", "通过"],
    ["TC-12", "成绩计算", "考勤 100/表现 80/实验 90/作业 100，权重 20/20/30/30", "平时 93.0，总成绩按占比", "通过"],
    ["TC-13", "删除保护", "删除仍有学生的班级", "409 并提示", "通过"],
    ["TC-14", "级联删除", "删除课程", "相关考勤/成绩/实验/作业清理", "通过"],
    ["TC-15", "成绩导出", "导出总成绩 CSV", "242 行，Excel 中文正常", "通过"],
    ["TC-16", "学生数据范围", "学生查询考勤", "仅返回本人记录", "通过"],
    ["TC-17", "教师数据范围", "教师查询考勤", "仅返回所授课程", "通过"],
    ["TC-18", "批量重复", "批量中混入已存在记录", "跳过并计数，不中断", "通过"],
    ["TC-19", "表现删除", "删除一条表现记录", "删除成功并刷新", "通过"],
    ["TC-20", "实验记录覆盖", "同一学生同一实验提交两次", "仅一条记录，分数更新", "通过"],
    ["TC-21", "作业记录覆盖", "同一学生同一作业提交两次", "仅一条记录，状态更新", "通过"],
    ["TC-22", "新增学生账号", "新增学生", "自动创建登录账号", "通过"],
    ["TC-23", "教师删除保护", "删除仍有课程的教师", "409 并提示", "通过"],
    ["TC-24", "课程级联删除", "删除课程", "考勤/表现/实验/作业/成绩/规则清理", "通过"],
    ["TC-25", "规则唯一", "同一课程重复新增规则", "数据库唯一约束拒绝", "通过"],
    ["TC-26", "密码强度", "新密码少于 6 位", "400 并提示", "通过"],
    ["TC-27", "会话失效", "改密后其他设备访问", "旧会话失效", "通过"],
    ["TC-28", "分页边界", "请求超出总页数", "自动回到最后一页", "通过"],
    ["TC-29", "CSV 编码", "Excel 打开导出文件", "中文不乱码", "通过"],
    ["TC-30", "审计记录", "执行增删改后查看日志", "记录操作人、动作、时间、IP", "通过"],
], widths=[1.6, 3.0, 6.4, 5.2, 2.0], font_size=9)

h2("7.11 系统运行效果")
body("""系统运行截图包括登录页、工作台、考勤录入、成绩汇总、基础数据管理与操作日志。截图中展示了前后端分离版的实际界面与交互效果，具体见下图。""")
shots = [
    ("login.png", "图 7-1 登录页面"),
    ("dashboard.png", "图 7-2 工作台"),
    ("attendance.png", "图 7-3 考勤查询与录入"),
    ("scores.png", "图 7-4 成绩汇总与导出"),
    ("students.png", "图 7-5 学生管理"),
    ("audit.png", "图 7-6 操作日志"),
]
for name, cap in shots:
    path = os.path.join(SHOT, name)
    if os.path.exists(path):
        image(path, 12.5, cap)
    else:
        p = doc.add_paragraph("【请在此处插入截图：" + cap[2:] + "】", style="CaptionCN")
        p.runs[0].font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)

h1("八．项目总结及收获")

h2("8.1 实训过程总结")
body("""本次实训从需求分析开始，经历了环境搭建、数据库设计、后端开发、前端开发、测试调试与容器化部署等阶段。开发初期先完成数据库表结构与演示数据，保证各功能模块有稳定的数据基础；随后按“登录认证—基础数据—教学过程—成绩统计”的顺序实现功能；最后补充前后端分离接口与 Docker 部署，使系统具备完整的演示与交付能力。""")
body("""在实现过程中，我逐步建立了“先设计、后编码、再验证”的工作习惯：每新增一个功能，先明确数据表、接口与权限，再编写代码，最后通过页面操作与接口请求双重验证。对于成绩计算、批量考勤、级联删除等容易出错的逻辑，均补充了边界校验与测试用例。""")

h2("8.2 技术能力收获")
body("""通过本实训，我较为系统地掌握了 Spring Boot 的分层开发方式，理解了 Controller、Service、Mapper 的职责划分；熟悉了 MyBatis 注解 SQL、动态 SQL、联合唯一约束与聚合查询；掌握了 Session 登录态、拦截器、BCrypt 密码存储、同源校验等安全实践；学会了使用外键与事务保证数据一致性。""")
body("""在工程能力方面，我掌握了 Maven 构建、WAR 打包、Tomcat 部署、Nginx 反向代理与 Docker Compose 编排；理解了前后端分离的接口设计、Cookie 同源策略与静态资源托管方式；也学会了通过日志、接口返回与数据库查询定位问题。""")

h2("8.3 遇到的问题及解决方法")
body("""（1）IDE 使用 JDK 25 编译，运行时报 UnsupportedClassVersionError（class 文件版本 69.0）。解决方法：将 IDE 编译 JDK 固定为 17，执行 mvn clean package 重新构建，并检查所有 class 均为版本 61。""")
body("""（2）应用启动时自动执行数据脚本，导致演示数据重复。解决方法：启动时只执行幂等的建表脚本，演示数据改为手工执行一次快照脚本。""")
body("""（3）Thymeleaf 3.1 不允许在 th:onclick 中直接拼接字符串，页面按钮失效。解决方法：改用 data-* 属性传递参数，由外部 JS 绑定事件。""")
body("""（4）MyBatis 中 LIKE 拼接写成 '%' + #{keyword} + '%' 在 MySQL 下语法错误。解决方法：改为 CONCAT('%', #{keyword}, '%')。""")
body("""（5）循环更新时出现 ResultSet closed。解决方法：先将查询结果物化为 List，再执行写操作。""")
body("""（6）成绩报表按学生逐条查询导致 N+1，页面响应慢。解决方法：改为每类数据一次聚合查询，再用 Map 匹配，响应时间从数秒降到亚秒级。""")
body("""（7）演示数据日期不真实，考勤几乎全部集中在周五。解决方法：按 16 个教学周重新生成数据，不同课程分布在不同工作日，避开周末与国庆假期。""")

h2("8.4 自身不足与改进方向")
body("""系统虽然完成了预定功能，但仍存在不足。第一，Service 层覆盖不完整，考勤、表现、实验、作业等模块的部分业务逻辑仍在 Controller 中，后续应按模块拆分 Service，使分层更加彻底，也便于单元测试。第二，数据库实际采用 MySQL，与指导手册中的 SQL Server 存在差异，虽然核心 SQL 兼容，但若课程要求严格，仍需补充 SQL Server 的建表脚本与驱动配置。第三，前后端分离版的功能仍在持续补齐，部分管理操作的历史交互习惯还需要优化，例如批量导入、表格排序与筛选条件的记忆。""")
body("""第四，测试以功能测试与少量单元测试为主，缺少自动化接口测试与性能测试；后续可以引入 MockMvc、Testcontainers 与 JMeter，对关键接口做回归与压力测试。第五，成绩导出目前为 CSV，若需要更正式的报表，可引入 Apache POI 生成带格式的 Excel，或使用模板引擎生成 PDF。第六，系统在高并发、缓存、日志采集、权限精细化（菜单级/按钮级）方面仍是空白，距离生产系统还有差距。""")
body("""针对上述不足，我计划从三个方面改进：一是继续完善 Service 分层与单元测试，提升代码可维护性；二是补充 SQL Server 兼容脚本与部署文档，满足不同环境要求；三是完善前端交互与报表能力，增加批量导入、图表统计与导出格式选择。通过持续迭代，使系统从“能运行的课程作业”逐步接近“可用于实际教学的业务系统”。""")

# 附录
h1("附录 A 数据库脚本说明")
body("""数据库脚本位于 src/main/resources/db/ 目录：schema-mysql.sql 负责建表、索引与外键，幂等可重复执行；data-mysql.sql 为完整演示数据快照，执行前会清空业务表；migration-fk.sql 用于已有库补充外键。执行顺序为：建库 → schema-mysql.sql → data-mysql.sql。""")
body("""数据快照包含 8 个班级、240 名学生、5 名教师、12 门课程、76 条课程班级关系、12 条评价规则，以及约 6.5 万条教学记录。所有日期均按 2025 年秋季学期 16 个教学周生成，工作日上课，避开国庆假期。""")

h1("附录 B 主要接口清单")
caption("表 B-1 主要 REST 接口")
table(["分类", "接口", "说明"], [
    ["认证", "POST /api/auth/login、/logout、/password", "登录、退出、改密"],
    ["查询", "GET /api/dashboard、/courses、/classes", "统计与基础列表"],
    ["查询", "GET /api/students、/attendance、/performance", "分页查询"],
    ["查询", "GET /api/experiments、/homework、/scores、/audit", "教学数据与审计"],
    ["管理", "POST/DELETE /api/admin/students、teachers、classes", "基础数据维护"],
    ["管理", "POST/DELETE /api/admin/courses、rules", "课程与规则维护"],
    ["教学", "POST /api/attendance、/attendance/batch", "考勤录入"],
    ["教学", "POST /api/performance、/experiments、/homework", "表现/实验/作业"],
    ["教学", "POST /api/exams", "期末成绩录入"],
    ["导出", "GET /api/scores/export", "成绩 CSV 导出"],
], widths=[2.6, 9.0, 5.0], font_size=9.5)

h1("附录 C 运行说明")
body("""方式一（JAR/WAR 直接运行）：安装 JDK 17 与 MySQL 8，建库并执行 schema-mysql.sql 与 data-mysql.sql，执行 mvn clean package 生成 target/course-management-system-1.0.0.war，运行 java -jar target/course-management-system-1.0.0.war --server.port=1001，访问 http://localhost:1001。""")
body("""方式二（Docker Compose）：启动 Docker Desktop，在项目根目录执行 docker compose up -d --build，等待三个容器启动完成，访问 http://localhost:8080（前后端分离版）或 http://localhost:1001（Thymeleaf 版）。""")
body("""演示账号：管理员 admin/123456，教师 T1001/123456，学生 202201001/123456。首次使用默认密码登录后，系统会强制跳转修改密码页面，修改成功后使用新密码登录。""")

h1("参考文献")
body("""[1] Spring Boot Reference Documentation. https://docs.spring.io/spring-boot/.
[2] MyBatis 3 Reference Documentation. https://mybatis.org/mybatis-3/.
[3] MySQL 8.0 Reference Manual. https://dev.mysql.com/doc/refman/8.0/en/.
[4] Thymeleaf Documentation. https://www.thymeleaf.org/documentation.html.
[5] Docker Documentation. https://docs.docker.com/.""")

h1("附录 D 核心建表 SQL")
body("""以下为 schema-mysql.sql 中核心表的建表语句（节选），完整脚本见项目 src/main/resources/db/ 目录。""")
code("""CREATE TABLE IF NOT EXISTS student (
    id INT NOT NULL AUTO_INCREMENT,
    student_no VARCHAR(30) NOT NULL UNIQUE,
    name VARCHAR(50) NOT NULL,
    gender VARCHAR(4),
    clazz_id INT NOT NULL,
    phone VARCHAR(20),
    email VARCHAR(80),
    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
    PRIMARY KEY (id),
    KEY idx_student_clazz (clazz_id),
    CONSTRAINT fk_student_clazz FOREIGN KEY (clazz_id)
        REFERENCES clazz(id) ON DELETE RESTRICT
);

CREATE TABLE IF NOT EXISTS course (
    id INT NOT NULL AUTO_INCREMENT,
    course_no VARCHAR(30) NOT NULL UNIQUE,
    course_name VARCHAR(100) NOT NULL,
    credit DECIMAL(3,1) NOT NULL DEFAULT 2.0,
    semester VARCHAR(20) NOT NULL,
    teacher_id INT NOT NULL,
    PRIMARY KEY (id),
    CONSTRAINT fk_course_teacher FOREIGN KEY (teacher_id)
        REFERENCES teacher(id) ON DELETE RESTRICT
);

CREATE TABLE IF NOT EXISTS attendance (
    id INT NOT NULL AUTO_INCREMENT,
    student_id INT NOT NULL,
    course_id INT NOT NULL,
    attend_date DATE NOT NULL,
    status VARCHAR(20) NOT NULL,
    remark VARCHAR(200),
    recorder_id INT,
    UNIQUE (student_id, course_id, attend_date),
    PRIMARY KEY (id),
    CONSTRAINT fk_att_student FOREIGN KEY (student_id)
        REFERENCES student(id) ON DELETE CASCADE,
    CONSTRAINT fk_att_course FOREIGN KEY (course_id)
        REFERENCES course(id) ON DELETE CASCADE
);""")

# 成绩评定表放到最后
end_break = doc.add_paragraph()
end_break.paragraph_format.page_break_before = True
body_el.insert(len(body_el) - 1, table_el)

# 保存
os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print("saved:", OUT)
print("paragraphs:", len(doc.paragraphs), "tables:", len(doc.tables))
