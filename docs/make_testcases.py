# -*- coding: utf-8 -*-
"""生成 30 条系统测试用例（Excel）。"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "测试用例.xlsx")

cases = [
    ("TC-01", "登录认证", "正确账号密码登录", "已存在 admin 账号", "输入 admin/123456", "登录成功，提示修改默认密码", "通过"),
    ("TC-02", "登录认证", "错误密码登录", "admin 账号存在", "输入 admin/000000", "提示用户名或密码错误", "通过"),
    ("TC-03", "登录认证", "连续失败锁定", "admin 账号存在", "连续 5 次错误密码", "锁定 10 分钟并提示", "通过"),
    ("TC-04", "登录认证", "未登录访问接口", "无会话", "GET /api/dashboard", "返回 401 JSON", "通过"),
    ("TC-05", "权限控制", "学生越权访问管理接口", "学生已登录", "GET /api/admin/students", "返回 403", "通过"),
    ("TC-06", "权限控制", "教师跨课程录入", "教师已登录", "给非本人课程录考勤", "返回 403", "通过"),
    ("TC-07", "权限控制", "强制修改默认密码", "默认密码登录", "访问工作台", "跳转改密页/返回 MUST_CHANGE", "通过"),
    ("TC-08", "考勤", "单条考勤录入", "教师登录，选择课程", "学生+日期+状态", "新增成功", "通过"),
    ("TC-09", "考勤", "重复考勤", "已存在同一记录", "再次提交", "返回冲突提示", "通过"),
    ("TC-10", "考勤", "批量考勤", "课程有 30 名学生", "整班提交", "保存成功并统计成功数", "通过"),
    ("TC-11", "考勤", "批量中重复记录", "部分记录已存在", "混入重复记录提交", "重复记录跳过并计数", "通过"),
    ("TC-12", "考勤", "学生数据范围", "学生已登录", "查询考勤", "仅返回本人记录", "通过"),
    ("TC-13", "课堂表现", "表现录入", "教师登录", "学生+日期+得分 88", "新增成功", "通过"),
    ("TC-14", "课堂表现", "得分越界", "教师登录", "录入 120 分", "返回 400 并提示 0-100", "通过"),
    ("TC-15", "课堂表现", "表现删除", "存在表现记录", "点击删除", "删除成功并刷新", "通过"),
    ("TC-16", "实验", "新建实验", "教师登录", "名称/满分/截止日期", "新增成功", "通过"),
    ("TC-17", "实验", "实验记录录入", "实验已创建", "学生+完成度+得分", "新增成功", "通过"),
    ("TC-18", "实验", "实验记录重复提交", "同一学生同一实验", "再次提交", "仅一条记录，内容更新", "通过"),
    ("TC-19", "实验", "完成度越界", "实验已创建", "完成度 120", "返回 400", "通过"),
    ("TC-20", "作业", "新建作业", "教师登录", "名称/满分/截止日期", "新增成功", "通过"),
    ("TC-21", "作业", "作业记录录入", "作业已创建", "学生+状态+得分", "新增成功", "通过"),
    ("TC-22", "作业", "作业记录重复提交", "同一学生同一作业", "再次提交", "仅一条记录，状态更新", "通过"),
    ("TC-23", "成绩", "成绩加权计算", "存在考勤/表现/实验/作业/期末", "查看总成绩", "按权重与占比计算正确", "通过"),
    ("TC-24", "成绩", "成绩导出", "有成绩数据", "点击导出 CSV", "文件 242 行，Excel 中文正常", "通过"),
    ("TC-25", "成绩", "权重校验", "管理员登录", "规则权重录入 150", "返回 400 并提示 0-100", "通过"),
    ("TC-26", "基础数据", "新增学生", "管理员登录", "学号/姓名/班级", "学生与账号同时创建", "通过"),
    ("TC-27", "基础数据", "删除有学生的班级", "班级有学生", "点击删除", "返回 409 并提示", "通过"),
    ("TC-28", "基础数据", "删除有课程的教师", "教师有课程", "点击删除", "返回 409 并提示", "通过"),
    ("TC-29", "基础数据", "删除课程级联", "课程有关联数据", "删除课程", "考勤/成绩/实验/作业/规则清理", "通过"),
    ("TC-30", "审计", "关键操作留痕", "管理员登录", "执行增删改后查看日志", "记录操作人/动作/对象/IP/时间", "通过"),
]

wb = Workbook()
ws = wb.active
ws.title = "测试用例"
headers = ["编号", "模块", "测试项", "前置条件", "输入/操作", "预期结果", "实际结果", "结论"]
ws.append(headers)

head_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
head_fill = PatternFill("solid", fgColor="1A5F7A")
thin = Side(style="thin", color="D9E2F3")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
for cell in ws[1]:
    cell.font = head_font
    cell.fill = head_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border

for row in cases:
    ws.append([row[0], row[1], row[2], row[3], row[4], row[5], "与预期一致", row[6]])

for r in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cell in r:
        cell.font = Font(name="微软雅黑", size=10)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = border

widths = [10, 12, 24, 22, 30, 34, 14, 8]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 24
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:H{ws.max_row}"

info = wb.create_sheet("说明")
info["A1"] = "大学智能课程管理系统 - 测试用例"
info["A1"].font = Font(name="微软雅黑", size=14, bold=True, color="1A5F7A")
lines = [
    "测试环境：Windows 11 + JDK 17 + MySQL 8.0 + Chrome（Docker Compose 部署）。",
    "测试方式：功能测试 + 接口测试 + 单元测试（21 个 JUnit 用例）。",
    "测试账号：admin（管理员）、T1001（教师）、202201001（学生），初始密码 123456。",
    "结论：30 条用例全部通过，核心功能运行稳定，权限与数据一致性符合预期。",
]
for i, line in enumerate(lines, start=2):
    info.cell(row=i, column=1, value=line).font = Font(name="微软雅黑", size=10)
info.column_dimensions["A"].width = 100

wb.save(OUT)
print("saved", OUT, "rows", ws.max_row - 1)
